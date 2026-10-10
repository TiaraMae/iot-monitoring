"""Device pairing, unpairing, and node event handling."""

import json
import time
from collections import defaultdict
from datetime import datetime, timedelta

import requests

from flask import Blueprint, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app import config, models, mqtt
from app import calibration
from app.db import get_conn, release_conn


devices_bp = Blueprint("devices", __name__)

EVENT_DEDUPE_CACHE = {}


def _is_dryer(app_type):
    return "Dryer" in (app_type or "")


def _daily_export_rows(is_dry, raw, voltage):
    """Build (headers, rows, day_count) for the per-day aggregate Excel export.

    Rows cover RUNNING readings only. The sensor-value averages skip the
    first AVG_DELTA_T_WARMUP_MINUTES of each run (same rule as the on-screen
    HVAC Daily Averages table); Avg Current covers the WHOLE run (inrush
    included) so it stays consistent with the Energy column, which
    integrates the full running cycles of each day.

    A trailing "Total (all days)" row is appended when at least one day
    exists: averages-of-averages for the sensor values and Avg Current, and
    the sum of the per-day Energy values (energy accumulates per day).
    day_count excludes the total row.
    """
    if is_dry:
        dicts = [{"time": r[0], "texhaust": r[1], "rh_exhaust": r[2],
                  "pressure": r[3], "imotor": r[5]} for r in raw]
        fields = ["texhaust", "rh_exhaust", "pressure"]
        current_idx = 5
        headers = ["Date", "Avg Exhaust Temp (°C)", "Avg Exhaust RH (%)",
                   "Avg Gauge Pressure (hPa)", "Avg Current (A)",
                   "Energy (kWh)"]
    else:
        dicts = [{"time": r[0], "treturn": r[1], "tsupply": r[2],
                  "icompressor": r[3]} for r in raw]
        fields = ["treturn", "tsupply"]
        current_idx = 3
        headers = ["Date", "Avg Return Temp (°C)", "Avg Supply Temp (°C)",
                   "Avg Delta-T (°C)", "Avg Current (A)",
                   "Energy (kWh)"]
    stats = models.compute_daily_export_averages(
        dicts, config.RUNNING_CURRENT_THRESHOLD,
        config.AVG_DELTA_T_WARMUP_MINUTES, fields,
        current_key="imotor" if is_dry else "icompressor")
    by_date = defaultdict(list)
    for r in raw:
        by_date[r[0].date()].append((r[0], r[current_idx]))
    rows = []
    for s in reversed(stats):  # chronological order, oldest day first
        day = datetime.fromisoformat(s["date"]).date()
        energy = models.compute_daily_energy(by_date.get(day, []), voltage)
        if is_dry:
            rows.append([s["date"], s["avg_texhaust"], s["avg_rh_exhaust"],
                         s["avg_pressure"], s["avg_current"], energy])
        else:
            avg_delta = (round(s["avg_treturn"] - s["avg_tsupply"], 2)
                         if s["avg_treturn"] is not None and s["avg_tsupply"] is not None
                         else None)
            rows.append([s["date"], s["avg_treturn"], s["avg_tsupply"],
                         avg_delta, s["avg_current"], energy])
    day_count = len(rows)
    if rows:
        total = ["Total (all days)"]
        # All columns except Energy are averages -> average-of-averages;
        # Energy accumulates per day -> sum.
        for idx in range(1, len(headers) - 1):
            vals = [r[idx] for r in rows if r[idx] is not None]
            total.append(round(sum(vals) / len(vals), 3) if vals else None)
        vals = [r[len(headers) - 1] for r in rows if r[len(headers) - 1] is not None]
        total.append(round(sum(vals), 4) if vals else None)
        rows.append(total)
    return headers, rows, day_count


def get_cf_deductor(current_sensor):
    """CF/deductor for a current-sensor model, from the single source of truth
    in config. Unknown/legacy sensors fall back to the default profile."""
    profile = config.CURRENT_SENSOR_PROFILES.get(
        current_sensor, config.CURRENT_SENSOR_PROFILES[config.DEFAULT_CURRENT_SENSOR]
    )
    return profile["cf"], profile["deductor"]


def _dedupe_event(mac, event_type):
    key = f"{mac}_{event_type}"
    now = time.time()
    if key in EVENT_DEDUPE_CACHE:
        if now - EVENT_DEDUPE_CACHE[key] < 5.0:
            return True
    EVENT_DEDUPE_CACHE[key] = now
    return False


def _send_type_config(mac, app_type, current_sensor, status):
    cf, deductor = get_cf_deductor(current_sensor)
    mqtt.send_node_command(mac, f"settype:{'dryer' if _is_dryer(app_type) else 'hvac'}")
    mqtt.send_node_command(mac, f"setcf:{cf}")
    mqtt.send_node_command(mac, f"setdeductor:{deductor}")
    if _is_dryer(app_type) or status == "normal":
        mqtt.send_node_command(mac, "restore:normal")
    else:
        mqtt.send_node_command(mac, "restore:offsetcalibrationneeded")


def handle_node_event(mac, payload):
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        print(f"Invalid event payload from {mac}")
        return

    event_type = data.get("event")
    if not event_type:
        return

    # calibration_progress events repeat every ~2 s while calibrating; the
    # 5 s dedupe window would swallow most of them, so they are exempt.
    if event_type != "calibration_progress" and _dedupe_event(mac, event_type):
        print(f"Duplicate event {event_type} from {mac} ignored.")
        return

    node = models.get_sensor_node_by_mac(mac)

    if event_type == "event_request_config":
        if not node:
            models.register_unpaired_node(mac)
            print(f"Auto-registered unknown node {mac}")
            mqtt.send_node_command(mac, "settype:unpaired")
            return

        models.update_node_last_seen(node["id"])
        if node["appliance_id"] is None or node["status"] != "paired":
            mqtt.send_node_command(mac, "settype:unpaired")
            return

        appliance = models.get_appliance(node["appliance_id"])
        if not appliance:
            mqtt.send_node_command(mac, "settype:unpaired")
            return

        # Node (re)connected: if it died mid-calibration, the DB can be stuck at
        # `calibrating` forever (fail events are not buffered) — reconcile first.
        from app.calibration import reconcile_on_node_connect
        reconcile_on_node_connect(appliance["id"])
        appliance = models.get_appliance(node["appliance_id"])

        _send_type_config(
            mac,
            appliance["type"],
            appliance["current_sensor"],
            appliance["operational_status"],
        )
        return

    if event_type == "checkin":
        if node:
            models.update_node_last_seen(node["id"])
        return

    if not node or node["appliance_id"] is None or node["status"] != "paired":
        if event_type in ("event_button2_offset_calibration_request",):
            mqtt.send_node_command(mac, "actiondenied:busy")
        return

    appliance = models.get_appliance(node["appliance_id"])
    if not appliance:
        return

    if event_type == "maintenance_request":
        if _is_dryer(appliance["type"]) or appliance["operational_status"] == "normal":
            models.insert_sensor_event(mac, "maintenance")
            mqtt.send_node_command(mac, "maintenanceack")
        else:
            mqtt.send_node_command(mac, "maintenancedenied")
        return

    if event_type == "event_button2_offset_calibration_request":
        calibration.handle_offset_calibration_event(mac)
        return

    if event_type == "calibration_progress":
        calibration.handle_calibration_progress(node["appliance_id"], data)
        return

    if event_type == "calibration_success_request":
        calibration.handle_calibration_success_request(mac, appliance, data)
        return

    if event_type == "calibration_fail_request":
        calibration.handle_calibration_fail_request(mac, node["appliance_id"])
        return


@devices_bp.route("/devices/pair", methods=["POST"])
@login_required
def pair_device():
    name = request.form.get("name", "").strip()
    app_type = request.form.get("type", "").strip()
    current_sensor = request.form.get("current_sensor", "").strip()
    node_id = request.form.get("node_id", "").strip()
    is_inverter = request.form.get("is_inverter") in ("on", "true", "1")

    if not name or not app_type or not node_id:
        flash("All fields are required", "error")
        return redirect(url_for("dashboard"))

    if app_type not in config.VALID_DEVICE_TYPES:
        flash("Invalid appliance type", "error")
        return redirect(url_for("dashboard"))

    if current_sensor not in config.CURRENT_SENSOR_PROFILES:
        flash("Invalid current sensor", "error")
        return redirect(url_for("dashboard"))

    try:
        node_id = int(node_id)
    except ValueError:
        flash("Invalid node", "error")
        return redirect(url_for("dashboard"))

    node = models.get_node_by_id(node_id)
    if not node or node["status"] != "unpaired":
        flash("Node not available", "error")
        return redirect(url_for("dashboard"))

    cf, deductor = get_cf_deductor(current_sensor)
    initial_status = "normal" if _is_dryer(app_type) else "offset_calibration_needed"

    try:
        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO appliances
            (user_id, name, type, is_inverter, location, brand,
             operational_status, cf, deductor, current_sensor, initial_calibrated_at)
            VALUES (%s, %s, %s, %s, 'Home', 'Generic', %s, %s, %s, %s, %s)
            RETURNING id
        """, (current_user.id, name, app_type, is_inverter,
              initial_status, cf, deductor, current_sensor,
              # Dryers never calibrate: their display window starts at pairing.
              config.now_wib() if _is_dryer(app_type) else None))
        appliance_id = cur.fetchone()[0]
        cur.execute("""
            UPDATE sensor_nodes
            SET appliance_id = %s, status = 'paired'
            WHERE id = %s
            RETURNING mac_address
        """, (appliance_id, node_id))
        mac = cur.fetchone()[0]
        conn.commit()
        cur.close()
        release_conn(conn)

        _send_type_config(mac, app_type, current_sensor, initial_status)

        if _is_dryer(app_type):
            flash(f"{name} added and ready. Configure baseline to enable alerts.", "success")
        else:
            flash(f"{name} added. Hold the hidden button for 5s in still air, then start the AC to calibrate offsets.", "success")
    except Exception as e:
        flash(f"Error pairing device: {e}", "error")

    return redirect(url_for("dashboard"))


@devices_bp.route("/devices/<int:appliance_id>/forget", methods=["POST"])
@login_required
def forget_device(appliance_id):
    from app import alerts

    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        flash("Device not found", "error")
        return redirect(url_for("dashboard"))

    mac = models.unpair_node_by_appliance(appliance_id)
    if mac:
        mqtt.send_node_command(mac, "settype:unpaired")

    alerts.clear_appliance_trackers(appliance_id)
    from app.calibration import CALIBRATION_TRACKER
    CALIBRATION_TRACKER.pop(appliance_id, None)
    flash("Device forgotten. All of its readings, alerts, and history were deleted.", "success")
    return redirect(url_for("dashboard"))


@devices_bp.route("/api/unpaired_nodes")
@login_required
def api_unpaired_nodes():
    return jsonify(models.get_unpaired_nodes())


@devices_bp.route("/api/device/<int:appliance_id>/latest")
@login_required
def api_device_latest(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    alert_status = models.get_appliance_alert_status(appliance_id)
    from datetime import datetime

    # Converge a dead calibration session (node cut off mid-calibration): if the
    # session cannot be live, revert before reporting status.
    if appliance["operational_status"] == "calibrating":
        from app.calibration import reconcile_calibration_state
        if reconcile_calibration_state(appliance):
            appliance = models.get_appliance(appliance_id)

    node = models.get_node_by_appliance(appliance_id)
    last_seen = node["last_seen"] if node else None
    # last_seen is stored via NOW() (naive LOCAL wall time) — compare on the same
    # clock. Comparing against aware UTC double-counts the timezone offset and
    # delays offline detection by ~7 h in UTC+7.
    # The offline window must match how often the node talks in its current
    # state: calibrated nodes stream telemetry every 10 s (strict 120 s), but
    # a need-calibration node is silent by design and only checks in every
    # 10 minutes — 120 s would flap it offline between checkins.
    offline_timeout = (config.NEED_CALIBRATION_OFFLINE_TIMEOUT_SECONDS
                       if appliance["operational_status"] == "offset_calibration_needed"
                       else config.OFFLINE_TIMEOUT_SECONDS)
    now = datetime.now()
    is_offline = (
        (now - last_seen).total_seconds() > offline_timeout
        if last_seen else True
    )

    if _is_dryer(appliance["type"]):
        row = models.get_latest_dryer_reading(appliance_id)
        if not row:
            return jsonify({
                "type": appliance["type"],
                "status": appliance["operational_status"],
                "alert_status": alert_status,
                "running_status": "idle",
                "is_offline": is_offline,
                "has_data": False,
                "calibrated": appliance["calibrated_at"] is not None,
            })
        imotor = max(0.0, row["imotor"] or 0)
        running = imotor >= config.RUNNING_CURRENT_THRESHOLD
        return jsonify({
            "time": row["time"].isoformat(),
            "Texhaust": row["texhaust"],
            "RHexhaust": row["rh_exhaust"],
            "Pressure": round(row["pressure"], 2) if row["pressure"] is not None else None,
            "AbsPressure": round(row["abs_pressure"], 2) if row["abs_pressure"] is not None else None,
            "Imotor": imotor,
            "type": appliance["type"],
            "status": appliance["operational_status"],
            "alert_status": alert_status,
            "running_status": "running" if running else "idle",
            "is_offline": is_offline,
            "has_data": True,
            "calibrated": appliance["calibrated_at"] is not None,
        })

    # HVAC
    row = models.get_latest_hvac_reading(appliance_id)
    if not row:
        return jsonify({
            "type": appliance["type"],
            "status": appliance["operational_status"],
            "alert_status": alert_status,
            "running_status": "idle",
            "is_offline": is_offline,
            "has_data": False,
            "calibrated": appliance["calibrated_at"] is not None,
        })

    icomp = row["icompressor"] or 0
    running = icomp >= config.RUNNING_CURRENT_THRESHOLD
    delta_t = round((row["treturn"] or 0) - (row["tsupply"] or 0), 2)
    return jsonify({
        "time": row["time"].isoformat(),
        "Treturn": row["treturn"],
        "Tsupply": row["tsupply"],
        "DeltaT": delta_t,
        "AvgDeltaT24h": models.get_avg_delta_t_running_24h(appliance_id),
        "Icompressor": icomp,
        "type": appliance["type"],
        "status": appliance["operational_status"],
        "alert_status": alert_status,
        "running_status": "running" if running else "idle",
        "is_offline": is_offline,
        "has_data": True,
        "calibrated": appliance["calibrated_at"] is not None,
    })


@devices_bp.route("/api/device/<int:appliance_id>/settings", methods=["GET", "POST"])
@login_required
def api_device_settings(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    if request.method == "GET":
        return jsonify({
            "id": appliance["id"],
            "name": appliance["name"],
            "type": appliance["type"],
            "sub_type": appliance["sub_type"],
            "is_inverter": appliance["is_inverter"],
            "current_sensor": appliance["current_sensor"],
            "voltage": appliance["voltage"],
            "cf": appliance["cf"],
            "deductor": appliance["deductor"],
            "alert_enabled": appliance["alert_enabled"],
            "delta_t_lcl": appliance["delta_t_lcl"],
            "delta_t_delay_minutes": appliance["delta_t_delay_minutes"],
            "atmospheric_pressure": appliance["atmospheric_pressure"],
            "treturn_offset": appliance["treturn_offset"],
            "tsupply_offset": appliance["tsupply_offset"],
            "operational_status": appliance["operational_status"],
        })

    data = request.get_json() or {}
    updates = {}
    bool_fields = {"alert_enabled"}
    numeric_fields = {
        "delta_t_lcl",
        "delta_t_delay_minutes",
        "voltage",
        "atmospheric_pressure",
    }

    for field in bool_fields | numeric_fields:
        if field not in data:
            continue
        if field in bool_fields:
            updates[field] = bool(data[field])
        else:
            val = data[field]
            if val in (None, ""):
                updates[field] = None
            else:
                try:
                    updates[field] = (
                        int(val) if field == "delta_t_delay_minutes" else float(val)
                    )
                except (ValueError, TypeError):
                    return jsonify({"error": f"Invalid value for {field}"}), 400

    if "delta_t_lcl" in updates:
        lcl = updates["delta_t_lcl"]
        if lcl is not None and lcl < 0:
            return jsonify({"error": "delta_t_lcl must be >= 0"}), 400
    if "delta_t_delay_minutes" in updates:
        delay = updates["delta_t_delay_minutes"]
        if delay is not None and delay < 1:
            return jsonify({"error": "delta_t_delay_minutes must be >= 1"}), 400

    # current_sensor is fixed at pairing time and intentionally NOT editable
    # here — changing a sensor model means different hardware on the node.

    if not updates:
        return jsonify({"error": "No valid fields provided"}), 400

    if models.update_appliance_settings(appliance_id, updates):
        return jsonify({"success": True})
    return jsonify({"error": "Failed to update settings"}), 500


@devices_bp.route("/api/device/<int:appliance_id>/alerts")
@login_required
def api_device_alerts(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404
    return jsonify({"alerts": models.get_alerts_for_appliance(appliance_id)})


@devices_bp.route("/api/alerts/<int:alert_id>/acknowledge", methods=["POST"])
@login_required
def api_acknowledge_alert(alert_id):
    # Ownership check via appliance join
    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT a.user_id FROM alerts al
            JOIN appliances a ON al.appliance_id = a.id
            WHERE al.id = %s
        """, (alert_id,))
        row = cur.fetchone()
        if not row or row[0] != current_user.id:
            return jsonify({"error": "not found"}), 404
    finally:
        cur.close()
        release_conn(conn)

    if models.resolve_alert(alert_id):
        return jsonify({"success": True})
    return jsonify({"error": "Failed to resolve"}), 500


@devices_bp.route("/api/alerts/<int:alert_id>/resolve", methods=["POST"])
@login_required
def api_resolve_alert(alert_id):
    return api_acknowledge_alert(alert_id)


@devices_bp.route("/api/device/<int:appliance_id>/latest_n")
@login_required
def api_device_latest_n(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    limit = request.args.get("limit", 120, type=int)
    limit = max(1, min(config.MAX_CHART_POINTS, limit))
    start = request.args.get("start")
    end = request.args.get("end")
    filtered = request.args.get("filtered", "true").lower() != "false"

    time_filter = ""
    params = [appliance_id, limit]
    if start and end:
        time_filter = " AND r.time >= %s AND r.time <= %s "
        params = [appliance_id, start, end, limit]

    is_dry = _is_dryer(appliance["type"])
    current_filter = " AND r.imotor >= {}".format(config.RUNNING_CURRENT_THRESHOLD) if (filtered and is_dry) else ""
    current_filter_hvac = " AND r.icompressor >= {}".format(config.RUNNING_CURRENT_THRESHOLD) if (filtered and not is_dry) else ""

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        if is_dry:
            cur.execute(f"""
                SELECT r.time, r.texhaust, r.rh_exhaust, r.pressure, r.abs_pressure, r.imotor
                FROM dryer_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                WHERE sn.appliance_id = %s {time_filter}{current_filter}
                ORDER BY r.time DESC
                LIMIT %s
            """, tuple(params))
            rows = cur.fetchall()
            return jsonify([
                {
                    "time": r[0].isoformat(),
                    "Texhaust": r[1],
                    "RHexhaust": r[2],
                    "Pressure": r[3],
                    "AbsPressure": r[4],
                    "Imotor": r[5],
                }
                for r in rows
            ])

        cur.execute(f"""
            SELECT r.time, r.treturn, r.tsupply, r.icompressor
            FROM hvac_readings r
            JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
            WHERE sn.appliance_id = %s {time_filter}{current_filter_hvac}
            ORDER BY r.time DESC
            LIMIT %s
        """, tuple(params))
        rows = cur.fetchall()
        return jsonify([
            {
                "time": r[0].isoformat(),
                "Treturn": r[1],
                "Tsupply": r[2],
                "DeltaT": round(abs((r[1] or 0) - (r[2] or 0)), 2),
                "Icompressor": r[3],
                "Tcoil": None,
                "RHreturn": None,
                "RHsupply": None,
            }
            for r in rows
        ])
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


@devices_bp.route("/api/device/<int:appliance_id>/maintenance_logs")
@login_required
def api_device_maintenance_logs(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404
    return jsonify({"logs": models.get_maintenance_logs_for_appliance(appliance_id)})


@devices_bp.route("/api/device/<int:appliance_id>/export_excel")
@login_required
def api_export_excel(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    filtered = request.args.get("filtered", "true").lower() != "false"

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        query_params = [appliance_id]
        date_filter = ""
        if start_date:
            date_filter += " AND r.time >= %s"
            query_params.append(start_date)
        if end_date:
            try:
                end_dt = datetime.fromisoformat(end_date)
                end_date = (end_dt + timedelta(seconds=1)).isoformat()
            except ValueError:
                pass
            date_filter += " AND r.time <= %s"
            query_params.append(end_date)

        is_dry = _is_dryer(appliance["type"])
        granularity = request.args.get("granularity", "points")
        if granularity not in ("points", "daily"):
            granularity = "points"
        voltage = models.get_appliance_voltage(appliance_id)

        # The idle/running SQL filter only applies to per-point exports; the
        # daily aggregate is running-only by definition (see _daily_export_rows).
        apply_current_filter = filtered and granularity == "points"
        current_filter = " AND r.imotor >= {}".format(config.RUNNING_CURRENT_THRESHOLD) if (apply_current_filter and is_dry) else ""
        current_filter_hvac = " AND r.icompressor >= {}".format(config.RUNNING_CURRENT_THRESHOLD) if (apply_current_filter and not is_dry) else ""

        if is_dry:
            cur.execute(f"""
                SELECT r.time, r.texhaust, r.rh_exhaust, r.pressure, r.abs_pressure, r.imotor
                FROM dryer_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                WHERE sn.appliance_id = %s {date_filter}{current_filter}
                ORDER BY r.time ASC
            """, tuple(query_params))
            raw = cur.fetchall()
        else:
            cur.execute(f"""
                SELECT r.time, r.treturn, r.tsupply, r.icompressor
                FROM hvac_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                WHERE sn.appliance_id = %s {date_filter}{current_filter_hvac}
                ORDER BY r.time ASC
            """, tuple(query_params))
            raw = cur.fetchall()

        if granularity == "daily":
            headers, rows, day_count = _daily_export_rows(is_dry, raw, voltage)
        elif is_dry:
            headers = ["Timestamp", "Exhaust Temp (°C)", "Exhaust RH (%)", "Gauge Pressure (hPa)", "Raw Absolute Pressure (hPa)", "Current (A)"]
            rows = []
            for r in raw:
                rows.append([r[0], r[1], r[2], r[3], r[4], r[5]])
        else:
            headers = ["Timestamp", "Return Temp (°C)", "Supply Temp (°C)", "Current (A)", "Delta-T (°C)"]
            rows = []
            for r in raw:
                delta = round(abs((r[1] or 0) - (r[2] or 0)), 2) if (r[1] is not None and r[2] is not None) else None
                rows.append([r[0], r[1], r[2], r[3], delta])

        wb = Workbook()
        ws = wb.active
        ws.title = "Sensor Data"
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(row=1, column=1, value=f"Device: {appliance['name']}").font = Font(size=14, bold=True)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(row=2, column=1, value=f"Type: {appliance['type']}")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
        ws.cell(row=3, column=1, value=f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))
        count_label = "Days" if granularity == "daily" else "Data Points"
        if granularity != "daily":
            day_count = len(rows)
        ws.cell(row=4, column=1, value=f"{count_label}: {day_count}")

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(rows, start=7):
            for j, val in enumerate(row, start=1):
                if j == 1 and isinstance(val, datetime):
                    ws.cell(row=i, column=j, value=val.strftime("%Y-%m-%d %H:%M:%S"))
                elif isinstance(val, float):
                    # 4 decimals: kWh energy values lose meaning at 3 (0.1082 -> 0.108).
                    ws.cell(row=i, column=j, value=round(val, 4))
                else:
                    ws.cell(row=i, column=j, value=val)

        if not rows:
            ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=len(headers))
            ws.cell(row=7, column=1, value="No data available for selected range.")

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        # Build a descriptive filename that includes the range/filter info.
        safe_name = appliance["name"].replace(" ", "_").replace("/", "_").replace("\\", "_")
        if granularity == "daily":
            filter_suffix = "daily"
        else:
            filter_suffix = "filtered" if filtered else "unfiltered"

        def _fmt_dt_filename(iso_str):
            try:
                dt = datetime.fromisoformat(iso_str)
                return dt.strftime("%m%d%Y-%H%M%S")
            except Exception:
                return iso_str.replace(":", "").replace("-", "").replace(" ", "_")

        if start_date and end_date:
            safe_start = _fmt_dt_filename(start_date)
            safe_end = _fmt_dt_filename(end_date)
            download_name = f"{safe_name}_export_{safe_start}_to_{safe_end}_{filter_suffix}.xlsx"
        else:
            export_ts = datetime.now().strftime("%m%d%Y-%H%M%S")
            download_name = f"{safe_name}_export_{export_ts}_{filter_suffix}.xlsx"

        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=download_name,
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


@devices_bp.route("/api/appliances")
@login_required
def api_appliances():
    return jsonify(models.get_appliances_for_user(current_user.id))


def _mask_webhook_url(url):
    if not url:
        return ""
    if len(url) <= 16:
        return url
    return url[:6] + "..." + url[-12:]


@devices_bp.route("/api/user/discord_webhook", methods=["GET", "POST"])
@login_required
def api_user_discord_webhook():
    if request.method == "GET":
        webhook_url = models.get_user_webhook_by_user_id(current_user.id) or ""
        return jsonify({
            "webhook_url": webhook_url,
            "url": webhook_url,
            "masked": _mask_webhook_url(webhook_url),
        })

    data = request.get_json() or {}
    webhook_url = (data.get("webhook_url") or data.get("url") or "").strip()
    if webhook_url and not webhook_url.startswith(("http://", "https://")):
        return jsonify({"error": "Invalid webhook URL"}), 400

    if models.update_user_webhook(current_user.id, webhook_url):
        return jsonify({"success": True})
    return jsonify({"error": "Failed to save webhook"}), 500


@devices_bp.route("/api/user/discord_webhook/test", methods=["POST"])
@login_required
def api_user_discord_webhook_test():
    data = request.get_json() or {}
    webhook_url = (data.get("webhook_url") or data.get("url") or "").strip()
    if not webhook_url:
        return jsonify({"error": "Webhook URL is required"}), 400

    try:
        payload = {
            "embeds": [
                {
                    "title": "Test Alert",
                    "description": "This is a test Discord alert from IoT Monitoring.",
                    "color": 0x2563EB,
                    "timestamp": datetime.now(config.TIMEZONE).isoformat(),
                }
            ]
        }
        resp = requests.post(webhook_url, json=payload, timeout=5)
        if resp.status_code in (200, 204):
            return jsonify({"success": True, "message": "Test alert sent"})
        return jsonify({"error": f"Discord returned {resp.status_code}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@devices_bp.route("/api/send_command", methods=["POST"])
@login_required
def api_send_command():
    data = request.get_json() or {}
    mac = (data.get("mac") or "").strip()
    command = (data.get("command") or "").strip()
    if not mac or not command:
        return jsonify({"error": "mac and command are required"}), 400
    mqtt.send_node_command(mac, command)
    return jsonify({"success": True, "message": f"Command '{command}' sent to {mac}"})

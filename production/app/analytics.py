"""Energy and cycle analytics endpoints."""

from collections import defaultdict
from datetime import datetime, timedelta
from io import BytesIO

from flask import Blueprint, request, jsonify, send_file
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.db import get_conn, release_conn
from app import config, models, mqtt


analytics_bp = Blueprint("analytics", __name__)


def _compute_motor_baseline_median(motor_readings, filter_threshold=None):
    if not motor_readings:
        return 0.0
    readings = motor_readings
    if filter_threshold is not None:
        readings = [r for r in readings if r <= filter_threshold]
        if not readings:
            return 0.0
    sorted_r = sorted(readings)
    n = len(sorted_r)
    if n % 2 == 1:
        return sorted_r[n // 2]
    return (sorted_r[n // 2 - 1] + sorted_r[n // 2]) / 2.0


def _compute_energy_kwh(readings, voltage):
    """Compute total energy (kWh) from (time, current) readings."""
    energy_ws = 0.0
    for i in range(1, len(readings)):
        prev_time, prev_current = readings[i - 1]
        curr_time, curr_current = readings[i]
        gap = (curr_time - prev_time).total_seconds()
        if gap <= 120 and prev_current >= 0.25:
            energy_ws += prev_current * voltage * gap
    return round(energy_ws / 3_600_000, 4)


def _compute_daily_energy(readings, voltage):
    """Compute daily energy splitting cycles by >120s gaps or current below threshold."""
    energy_ws = 0.0
    in_cycle = False
    cycle_readings = []
    for i, r in enumerate(readings):
        time_val, current = r
        current = float(current) if current is not None else 0.0
        if in_cycle and i > 0:
            gap = (time_val - readings[i - 1][0]).total_seconds()
            if gap > 120:
                for j in range(1, len(cycle_readings)):
                    dt = (cycle_readings[j][0] - cycle_readings[j - 1][0]).total_seconds()
                    energy_ws += cycle_readings[j - 1][1] * voltage * dt
                in_cycle = False
                cycle_readings = []
        if current >= 0.25 and not in_cycle:
            in_cycle = True
            cycle_readings = [(time_val, current)]
        elif in_cycle:
            if cycle_readings and cycle_readings[-1][0] != time_val:
                cycle_readings.append((time_val, current))
        if current < 0.25 and in_cycle:
            for j in range(1, len(cycle_readings)):
                dt = (cycle_readings[j][0] - cycle_readings[j - 1][0]).total_seconds()
                energy_ws += cycle_readings[j - 1][1] * voltage * dt
            in_cycle = False
            cycle_readings = []
    if in_cycle and cycle_readings:
        for j in range(1, len(cycle_readings)):
            dt = (cycle_readings[j][0] - cycle_readings[j - 1][0]).total_seconds()
            energy_ws += cycle_readings[j - 1][1] * voltage * dt
    return round(energy_ws / 3_600_000, 4)


@analytics_bp.route("/api/device/<int:appliance_id>/hvac_analytics")
@login_required
def hvac_analytics(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404
    if "Dryer" in appliance["type"]:
        return jsonify({"error": "Not HVAC"}), 400

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        start = request.args.get("start")
        end = request.args.get("end")
        if end:
            try:
                end_dt = datetime.fromisoformat(end)
                end = (end_dt + timedelta(seconds=1)).isoformat()
            except ValueError:
                pass

        if start and end:
            cur.execute("""
                SELECT DATE(r.time) as date,
                       AVG(r.treturn), AVG(r.tsupply)
                FROM hvac_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.icompressor >= 0.25
                  AND r.time >= a.created_at AND r.time >= %s AND r.time <= %s
                GROUP BY DATE(r.time)
                ORDER BY DATE(r.time) DESC
                LIMIT 30
            """, (appliance_id, start, end))
        else:
            cur.execute("""
                SELECT DATE(r.time) as date,
                       AVG(r.treturn), AVG(r.tsupply)
                FROM hvac_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.icompressor >= 0.25
                  AND r.time >= a.created_at
                GROUP BY DATE(r.time)
                ORDER BY DATE(r.time) DESC
                LIMIT 30
            """, (appliance_id,))
        daily_rows = cur.fetchall()
        daily_averages = [
            {
                "date": r[0].isoformat(),
                "avg_return": round(r[1], 2) if r[1] is not None else None,
                "avg_supply": round(r[2], 2) if r[2] is not None else None,
                "avg_intake": round(r[1], 2) if r[1] is not None else None,
                "avg_exit": round(r[2], 2) if r[2] is not None else None,
                "avg_coil": None,
            }
            for r in daily_rows
        ]

        voltage = models.get_appliance_voltage(appliance_id)
        if start and end:
            cur.execute("""
                SELECT r.time, r.icompressor
                FROM hvac_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.time >= a.created_at
                  AND r.time >= %s AND r.time <= %s
                ORDER BY r.time ASC
            """, (appliance_id, start, end))
        else:
            cur.execute("""
                SELECT r.time, r.icompressor
                FROM hvac_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.time >= a.created_at
                ORDER BY r.time ASC
            """, (appliance_id,))
        readings = cur.fetchall()

        readings_by_date = defaultdict(list)
        for r in readings:
            readings_by_date[r[0].date()].append(r)

        for day in daily_averages:
            date_key = datetime.fromisoformat(day["date"]).date()
            day_readings = readings_by_date.get(date_key, [])
            day["daily_energy_kwh"] = _compute_daily_energy(day_readings, voltage)

        return jsonify({"daily_averages": daily_averages})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


@analytics_bp.route("/api/device/<int:appliance_id>/dryer_analytics")
@login_required
def dryer_analytics(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404
    if "Dryer" not in appliance["type"]:
        return jsonify({"error": "Not a dryer"}), 400

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        cycle_start = 0.25
        baselines = models.get_spc_baselines(appliance_id)
        mean_current = baselines.get("current", {}).get("mean", 2.0)
        prominence_threshold = 0.40
        voltage = models.get_appliance_voltage(appliance_id)

        start = request.args.get("start")
        end = request.args.get("end")
        if end:
            try:
                end_dt = datetime.fromisoformat(end)
                end = (end_dt + timedelta(seconds=1)).isoformat()
            except ValueError:
                pass

        if start and end:
            cur.execute("""
                SELECT r.time, r.texhaust, r.rh_exhaust, r.pressure, r.imotor, r.abs_pressure
                FROM dryer_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.time >= a.created_at
                  AND r.time >= %s AND r.time <= %s
                ORDER BY r.time ASC
            """, (appliance_id, start, end))
        else:
            cur.execute("""
                SELECT r.time, r.texhaust, r.rh_exhaust, r.pressure, r.imotor, r.abs_pressure
                FROM dryer_readings r
                JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                JOIN appliances a ON a.id = sn.appliance_id
                WHERE sn.appliance_id = %s AND r.time >= a.created_at
                ORDER BY r.time ASC
            """, (appliance_id,))
        readings = cur.fetchall()
        if not readings:
            return jsonify([])

        cycles = []
        in_cycle = False
        current_cycle = {}
        _prev_current = 0.0
        _peak_state = "IDLE"
        _peak_max = 0.0
        _peak_valley = 0.0
        _peak_values = []
        _idle_start_time = None

        def _confirm_peak():
            nonlocal _peak_state, _peak_max, _peak_valley
            if _peak_state in ("RISING", "FALLING") and _peak_max > 0:
                prominence = _peak_max - _peak_valley
                if prominence >= prominence_threshold and _peak_max > mean_current + 0.15:
                    _peak_values.append(_peak_max)
            _peak_state = "IDLE"
            _peak_max = 0.0
            _peak_valley = 0.0

        for i, r in enumerate(readings):
            time_val, tex, rhex, press, imotor, _abs_press = r
            imotor = float(imotor) if imotor is not None else 0.0
            tex = float(tex) if tex is not None else 0.0
            rhex = float(rhex) if rhex is not None else None

            if in_cycle and i > 0:
                gap = (time_val - readings[i - 1][0]).total_seconds()
                if gap > 120:
                    in_cycle = False
                    prev_time = readings[i - 1][0]
                    _finalize_cycle_record(
                        current_cycle, prev_time, _peak_values, voltage, cycles
                    )
                    current_cycle = {}
                    _prev_current = 0.0
                    _peak_values = []
                    _idle_start_time = None

            if imotor >= cycle_start and not in_cycle:
                in_cycle = True
                current_cycle = {
                    "start_time": time_val,
                    "min_temp": tex,
                    "max_temp": tex,
                    "_rh_history": [],
                    "_currents": [],
                    "_times": [],
                    "_motor_readings": [],
                    "_pressures": [],
                }
                _prev_current = imotor
                _idle_start_time = None
                _confirm_peak()
                _peak_values = []

            if in_cycle:
                current_cycle["max_temp"] = max(current_cycle["max_temp"], tex)
                current_cycle["min_temp"] = min(current_cycle["min_temp"], tex)
                if rhex is not None:
                    current_cycle["_rh_history"].append(rhex)
                current_cycle["_currents"].append(imotor)
                current_cycle["_times"].append(time_val)
                if press is not None:
                    current_cycle["_pressures"].append(float(press))

                if imotor >= cycle_start:
                    current_cycle["_motor_readings"].append(imotor)
                    if imotor > _prev_current:
                        if _peak_state in ("IDLE", "FALLING"):
                            if _peak_state == "FALLING":
                                _confirm_peak()
                            _peak_valley = _prev_current
                        _peak_state = "RISING"
                        if imotor > _peak_max:
                            _peak_max = imotor
                    elif imotor < _prev_current:
                        if _peak_state == "RISING" and (_peak_max <= 0.1 or imotor < _peak_max - 0.1):
                            _peak_state = "FALLING"
                    _prev_current = imotor
                    _idle_start_time = None
                else:
                    if _idle_start_time is None:
                        _idle_start_time = time_val
                    idle_duration = (time_val - _idle_start_time).total_seconds()
                    if idle_duration >= 120:
                        in_cycle = False
                        current_cycle["end_time"] = readings[i - 1][0]
                        _finalize_cycle_record(
                            current_cycle, readings[i - 1][0], _peak_values, voltage, cycles
                        )
                        current_cycle = {}
                        _prev_current = 0.0
                        _peak_values = []
                        _idle_start_time = None

        if in_cycle:
            last_r = readings[-1]
            current_cycle["end_time"] = last_r[0]
            _finalize_cycle_record(
                current_cycle, last_r[0], _peak_values, voltage, cycles
            )

        cycles = [c for c in cycles if c.get("duration_minutes", 0) >= 1.0]
        cycles.reverse()
        return jsonify(cycles)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


def _finalize_cycle_record(current_cycle, end_time, peak_values, voltage, cycles):
    if not current_cycle or "start_time" not in current_cycle:
        return
    duration = end_time - current_cycle["start_time"]
    current_cycle["duration_minutes"] = round(duration.total_seconds() / 60, 1)
    valid_rh = [v for v in current_cycle.get("_rh_history", []) if v is not None]
    first_6 = valid_rh[:6] if len(valid_rh) > 6 else valid_rh
    current_cycle["start_rh"] = round(sum(first_6) / len(first_6), 2) if first_6 else None
    last_6 = valid_rh[-6:] if len(valid_rh) > 6 else valid_rh
    current_cycle["end_rh_avg"] = (
        round(sum(last_6) / len(last_6), 2) if last_6 else current_cycle.get("start_rh")
    )
    currents = current_cycle["_currents"]
    times = current_cycle["_times"]
    energy_ws = 0.0
    for j in range(1, len(currents)):
        dt = (times[j] - times[j - 1]).total_seconds()
        energy_ws += currents[j - 1] * voltage * dt
    current_cycle["energy_kwh"] = round(energy_ws / 3_600_000, 4)
    current_cycle["current_spike_avg"] = (
        round(sum(peak_values) / len(peak_values), 2) if peak_values else 0.0
    )
    current_cycle["ignition_count"] = len(peak_values)
    motor_readings = current_cycle.get("_motor_readings", [])
    filter_threshold = None
    if motor_readings:
        filter_threshold = (sum(motor_readings) / len(motor_readings)) * 1.15
    median_val = _compute_motor_baseline_median(motor_readings, filter_threshold=filter_threshold)
    current_cycle["motor_baseline_median"] = round(median_val, 3)
    pressures = [p for p in current_cycle.get("_pressures", []) if p is not None]
    current_cycle["avg_pressure"] = round(sum(pressures) / len(pressures), 2) if pressures else None

    for key in ("_rh_history", "_currents", "_times", "_motor_readings", "_pressures"):
        current_cycle.pop(key, None)
    current_cycle["start_time"] = current_cycle["start_time"].isoformat()
    current_cycle["end_time"] = end_time.isoformat()
    cycles.append(current_cycle)


HVAC_METRICS = ["deltat", "current"]
DRYER_METRICS = ["texhaust", "rhexhaust", "current", "pressure"]


@analytics_bp.route("/api/device/<int:appliance_id>/baseline_config", methods=["GET", "POST", "DELETE"])
@login_required
def api_baseline_config(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    dev_type = appliance["type"]
    expected = DRYER_METRICS if "Dryer" in dev_type else HVAC_METRICS

    if request.method == "GET":
        baselines = models.get_spc_baselines(appliance_id)
        result = {}
        for m in expected:
            if m in baselines:
                result[m] = baselines[m]
            else:
                result[m] = {"ucl": "", "lcl": "", "mean": ""}
        updated_at = models.get_spc_baselines_updated_at(appliance_id)
        return jsonify({
            "type": dev_type,
            "metrics": result,
            "updated_at": updated_at.isoformat() if updated_at else None,
            "delta_t_lcl": appliance.get("delta_t_lcl"),
            "delta_t_delay_minutes": appliance.get("delta_t_delay_minutes"),
        })

    if request.method == "DELETE":
        conn = get_conn()
        if not conn:
            return jsonify({"error": "db"}), 500
        cur = conn.cursor()
        try:
            cur.execute(
                "DELETE FROM spc_manual_baselines WHERE appliance_id = %s",
                (appliance_id,),
            )
            cur.execute(
                "UPDATE appliances SET baseline_configured = FALSE WHERE id = %s",
                (appliance_id,),
            )
            if "Dryer" not in dev_type:
                cur.execute(
                    "UPDATE appliances SET delta_t_lcl = NULL, "
                    "delta_t_delay_minutes = NULL WHERE id = %s",
                    (appliance_id,),
                )
            conn.commit()
            return jsonify({"success": True, "message": "Baseline removed"})
        except Exception as e:
            conn.rollback()
            return jsonify({"error": str(e)}), 500
        finally:
            cur.close()
            release_conn(conn)

    data = request.get_json() or {}
    metrics = data.get("metrics", {})
    baselines_to_save = {}
    for m in expected:
        if m not in metrics:
            continue
        try:
            ucl_raw = metrics[m].get("ucl")
            lcl_raw = metrics[m].get("lcl")
            mean_raw = metrics[m].get("mean")
            if "Dryer" in dev_type and m == "current" and mean_raw:
                mean_val = float(mean_raw)
                if (not ucl_raw or str(ucl_raw).strip() == "") and (
                    not lcl_raw or str(lcl_raw).strip() == ""
                ):
                    ucl = round(mean_val * 1.20, 3)
                    lcl = round(mean_val * 0.80, 3)
                else:
                    ucl = float(ucl_raw)
                    lcl = float(lcl_raw)
                if ucl <= lcl:
                    return jsonify({"error": f"UCL must be greater than LCL for {m}"}), 400
                baselines_to_save[m] = {"ucl": ucl, "lcl": lcl}
            elif "Dryer" in dev_type and m == "pressure":
                ucl = float(ucl_raw) if ucl_raw not in (None, "") else None
                if ucl is None:
                    return jsonify({"error": f"UCL is required for {m}"}), 400
                mean = float(mean_raw) if mean_raw not in (None, "") else None
                baselines_to_save[m] = {"mean": mean, "ucl": ucl, "lcl": None}
            else:
                ucl_str = str(ucl_raw).strip() if ucl_raw is not None else ""
                lcl_str = str(lcl_raw).strip() if lcl_raw is not None else ""
                if ucl_str == "" and lcl_str == "":
                    continue
                ucl = float(ucl_str) if ucl_str != "" else None
                lcl = float(lcl_str) if lcl_str != "" else None
                if ucl is not None and lcl is not None and ucl <= lcl:
                    return jsonify({"error": f"UCL must be greater than LCL for {m}"}), 400
                entry = {}
                if ucl is not None:
                    entry["ucl"] = ucl
                if lcl is not None:
                    entry["lcl"] = lcl
                if mean_raw not in (None, ""):
                    entry["mean"] = float(mean_raw)
                baselines_to_save[m] = entry
        except (ValueError, TypeError):
            return jsonify({"error": f"Invalid values for {m}"}), 400

    if not baselines_to_save:
        return jsonify({"error": "No valid baseline data provided"}), 400

    # Keep the alert engine's threshold store in sync: when a delta-T LCL is
    # saved for an HVAC unit, mirror it (plus the alert delay) onto the
    # appliances columns that check_hvac_delta_t_alert actually reads.
    alert_delay_minutes = None
    if "Dryer" not in dev_type and baselines_to_save.get("deltat", {}).get("lcl") is not None:
        raw_delay = data.get("alert_delay_minutes")
        if raw_delay not in (None, ""):
            try:
                alert_delay_minutes = int(float(raw_delay))
            except (ValueError, TypeError):
                return jsonify({"error": "alert_delay_minutes must be a number"}), 400
            if alert_delay_minutes < 1:
                return jsonify(
                    {"error": "alert_delay_minutes must be >= 1"}
                ), 400

    success, msg = models.save_spc_baselines(appliance_id, baselines_to_save)
    if success:
        if alert_delay_minutes is not None or (
            "Dryer" not in dev_type
            and baselines_to_save.get("deltat", {}).get("lcl") is not None
        ):
            delay = alert_delay_minutes
            if delay is None:
                delay = appliance.get("delta_t_delay_minutes")
            if delay is None:
                delay = config.DEFAULT_DELTA_T_DELAY_MINUTES
            conn = get_conn()
            if conn:
                cur = conn.cursor()
                try:
                    cur.execute(
                        "UPDATE appliances SET delta_t_lcl = %s, "
                        "delta_t_delay_minutes = %s WHERE id = %s",
                        (
                            baselines_to_save["deltat"]["lcl"],
                            delay,
                            appliance_id,
                        ),
                    )
                    conn.commit()
                except Exception as e:
                    conn.rollback()
                    return jsonify({"error": str(e)}), 500
                finally:
                    cur.close()
                    release_conn(conn)
        node = models.get_node_by_appliance(appliance_id)
        if node and node.get("mac_address"):
            mqtt.send_node_command(node["mac_address"], "baseline:set")
        return jsonify({"success": True, "message": msg})
    return jsonify({"error": msg}), 500


@analytics_bp.route("/api/device/<int:appliance_id>/spc_limits")
@login_required
def api_device_spc_limits(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    alert_status = models.get_appliance_alert_status(appliance_id)
    baselines = models.get_spc_baselines(appliance_id)
    expected = (
        ["texhaust", "rhexhaust", "pressure", "current"]
        if "Dryer" in appliance["type"]
        else ["deltat", "current"]
    )

    result = {
        "type": appliance["type"],
        "subtype": appliance["sub_type"],
        "status": appliance["operational_status"],
        "alert_status": alert_status,
        "baseline_configured": bool(appliance.get("baseline_configured")),
    }
    for metric in expected:
        b = baselines.get(metric, {})
        result[metric] = {
            "mean": b.get("mean"),
            "ucl": b.get("ucl"),
            "lcl": b.get("lcl"),
        }
    return jsonify(result)


@analytics_bp.route("/api/device/<int:appliance_id>/baseline_analysis")
@login_required
def api_device_baseline_analysis(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            """
            SELECT COUNT(*), MAX(updated_at)
            FROM spc_manual_baselines
            WHERE appliance_id = %s
            """,
            (appliance_id,),
        )
        count, updated_at = cur.fetchone()
        return jsonify({
            "type": appliance["type"],
            "status": appliance["operational_status"],
            "baseline_configured": count > 0,
            "baseline_set_at": updated_at.isoformat() if updated_at else None,
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        cur.close()
        release_conn(conn)


@analytics_bp.route("/api/device/<int:appliance_id>/pressure_baseline", methods=["GET", "POST"])
@login_required
def api_device_pressure_baseline(appliance_id):
    appliance = models.get_appliance(appliance_id)
    if not appliance or appliance["user_id"] != current_user.id:
        return jsonify({"error": "not found"}), 404

    if request.method == "GET":
        updated_at = appliance.get("updated_at")
        return jsonify({
            "atmospheric_pressure": appliance.get("atmospheric_pressure"),
            "updated_at": updated_at.isoformat() if updated_at else None,
        })

    data = request.get_json() or {}
    pressure = data.get("atmospheric_pressure")
    if pressure in (None, ""):
        return jsonify({"error": "atmospheric_pressure is required"}), 400
    try:
        pressure = float(pressure)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid atmospheric_pressure"}), 400

    if models.update_appliance_settings(appliance_id, {"atmospheric_pressure": pressure}):
        return jsonify({"success": True})
    return jsonify({"error": "Failed to update pressure baseline"}), 500


@analytics_bp.route("/api/energy_summary")
@login_required
def api_energy_summary():
    month_str = request.args.get("month", datetime.now().strftime("%Y-%m"))
    try:
        year, month = map(int, month_str.split("-"))
        month_start = datetime(year, month, 1)
        if month == 12:
            month_end = datetime(year + 1, 1, 1)
        else:
            month_end = datetime(year, month + 1, 1)
    except ValueError:
        return jsonify({"error": "Invalid month format. Use YYYY-MM."}), 400

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT id, name, type, created_at
            FROM appliances
            WHERE user_id = %s
            ORDER BY name
        """, (current_user.id,))
        appliances = cur.fetchall()

        result = []
        by_type = defaultdict(float)
        total_kwh = 0.0

        for app_id, app_name, app_type, created_at in appliances:
            voltage = models.get_appliance_voltage(app_id)
            query_start = max(month_start, created_at) if created_at else month_start
            if "Dryer" in app_type:
                cur.execute("""
                    SELECT r.time, r.imotor
                    FROM dryer_readings r
                    JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                    WHERE sn.appliance_id = %s AND r.time >= %s AND r.time < %s
                    ORDER BY r.time ASC
                """, (app_id, query_start, month_end))
            else:
                cur.execute("""
                    SELECT r.time, r.icompressor
                    FROM hvac_readings r
                    JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                    WHERE sn.appliance_id = %s AND r.time >= %s AND r.time < %s
                    ORDER BY r.time ASC
                """, (app_id, query_start, month_end))
            readings = cur.fetchall()
            energy = _compute_energy_kwh(readings, voltage)
            result.append({
                "id": app_id,
                "name": app_name,
                "type": app_type,
                "energy_kwh": energy,
            })
            by_type[app_type] += energy
            total_kwh += energy

        return jsonify({
            "month": month_str,
            "appliances": result,
            "total_kwh": round(total_kwh, 4),
            "by_type": dict(by_type),
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


@analytics_bp.route("/api/energy_summary/export")
@login_required
def api_energy_summary_export():
    month_str = request.args.get("month", datetime.now().strftime("%Y-%m"))
    try:
        year, month = map(int, month_str.split("-"))
        month_start = datetime(year, month, 1)
        if month == 12:
            month_end = datetime(year + 1, 1, 1)
        else:
            month_end = datetime(year, month + 1, 1)
    except ValueError:
        return jsonify({"error": "Invalid month format. Use YYYY-MM."}), 400

    conn = get_conn()
    if not conn:
        return jsonify({"error": "db"}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT id, name, type, created_at
            FROM appliances
            WHERE user_id = %s
            ORDER BY type, name
        """, (current_user.id,))
        appliances = cur.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = "Energy Summary"
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        ws.merge_cells("A1:E1")
        ws["A1"] = "Monthly Energy Consumption Summary"
        ws["A1"].font = Font(size=14, bold=True)
        ws.merge_cells("A2:E2")
        ws["A2"] = f"Month: {month_str}"
        ws["A2"].font = Font(size=11, bold=True)
        ws.merge_cells("A3:E3")
        ws["A3"] = f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        headers = ["Appliance Type", "Appliance Name", "Energy Consumption (kWh)"]
        row_idx = 5
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        total_kwh = 0.0
        for app_id, app_name, app_type, created_at in appliances:
            voltage = models.get_appliance_voltage(app_id)
            query_start = max(month_start, created_at) if created_at else month_start
            if "Dryer" in app_type:
                cur.execute("""
                    SELECT r.time, r.imotor
                    FROM dryer_readings r
                    JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                    WHERE sn.appliance_id = %s AND r.time >= %s AND r.time < %s
                    ORDER BY r.time ASC
                """, (app_id, query_start, month_end))
            else:
                cur.execute("""
                    SELECT r.time, r.icompressor
                    FROM hvac_readings r
                    JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
                    WHERE sn.appliance_id = %s AND r.time >= %s AND r.time < %s
                    ORDER BY r.time ASC
                """, (app_id, query_start, month_end))
            readings = cur.fetchall()
            energy = _compute_energy_kwh(readings, voltage)
            row_idx += 1
            ws.cell(row=row_idx, column=1, value=app_type)
            ws.cell(row=row_idx, column=2, value=app_name)
            ws.cell(row=row_idx, column=3, value=energy)
            total_kwh += energy

        row_idx += 1
        ws.cell(row=row_idx, column=1, value="Total")
        ws.cell(row=row_idx, column=1).font = Font(bold=True)
        ws.cell(row=row_idx, column=3, value=round(total_kwh, 4))
        ws.cell(row=row_idx, column=3).font = Font(bold=True)

        for col_idx in range(1, 4):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for r in range(1, ws.max_row + 1):
                cell_value = ws.cell(row=r, column=col_idx).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            ws.column_dimensions[col_letter].width = max_length + 2

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"Energy_Summary_{month_str}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)


@analytics_bp.route("/api/energy_months")
@login_required
def api_energy_months():
    conn = get_conn()
    if not conn:
        return jsonify([]), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT DISTINCT TO_CHAR(DATE_TRUNC('month', r.time), 'YYYY-MM') AS month
            FROM hvac_readings r
            JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
            JOIN appliances a ON a.id = sn.appliance_id
            WHERE a.user_id = %s
            UNION
            SELECT DISTINCT TO_CHAR(DATE_TRUNC('month', r.time), 'YYYY-MM') AS month
            FROM dryer_readings r
            JOIN sensor_nodes sn ON r.sensor_node_id = sn.id
            JOIN appliances a ON a.id = sn.appliance_id
            WHERE a.user_id = %s
            ORDER BY month DESC
        """, (current_user.id, current_user.id))
        rows = cur.fetchall()
        return jsonify([r[0] for r in rows])
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            cur.close()
            release_conn(conn)
